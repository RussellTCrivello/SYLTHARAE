"""Step 15 against PostgreSQL: report artifacts and manifests.

Executed on the migrated schema (m0023) with real runs. Properties:

* every format is rendered from the run's stored datasets; the stored SHA-256
  and size describe the stored bytes (checked by PostgreSQL itself) and the
  manifest names them, the run, its snapshot, fingerprints, counts and
  truncation; ``manifest_sha256`` is the digest of the canonical manifest;
* rendering is deterministic: the same run renders to the same bytes, later
  and for another user; a repeated request returns the existing artifact;
* artifacts are immutable and a row whose digest does not describe its bytes
  cannot be inserted; unsafe filenames cannot be stored;
* a tampered manifest is detected and the file is not handed out;
* visibility follows the run; viewers and deactivated creators are refused;
  only completed runs are rendered; unavailable formats are refused;
* hostile text stays text (HTML escaped, spreadsheet formulas guarded) and
  non-Latin text survives every format.
"""

from __future__ import annotations

import csv
import datetime
import hashlib
import io
import json
import time
import uuid
import zipfile
from types import SimpleNamespace

import psycopg2
import pytest

from core.reporting import REGISTRY
from core.reporting import render as renderers
from services.reporting import artifacts, runs

from _seed import connect, document, side, source

pytestmark = pytest.mark.integration

_U = uuid.uuid4().hex[:8]
HOSTILE = "=HYPERLINK(\"http://x\")<script>alert(1)</script>"
ARABIC = "تقرير_بغداد"


def _user(cur, role, name):
    cur.execute("INSERT INTO users (username, password_hash, role) VALUES (%s, 'x', %s)"
                " RETURNING id", (f"{name}_{_U}", role))
    uid = cur.fetchone()[0]
    return SimpleNamespace(id=uid, role=role, username=f"{name}_{_U}",
                           has_role=lambda *roles: role in roles)


@pytest.fixture(scope="module")
def world(pg_db, app):
    word = f"zart{_U}"
    conn = connect(pg_db)
    with conn, conn.cursor() as cur:
        s1, d1 = source(cur), side(cur)
        document(cur, source_id=s1, side_id=d1, text=f"art {word} one", file_type="pdf",
                 file_name=f"{HOSTILE}.pdf", file_date=datetime.date(2026, 3, 1))
        document(cur, source_id=s1, side_id=d1, text=f"art {word} two", file_type="pdf",
                 file_name=f"{ARABIC}.pdf", file_date=datetime.date(2026, 3, 2))
        document(cur, source_id=s1, side_id=d1, text=f"art {word} three", file_type="txt",
                 file_date=datetime.date(2026, 3, 3))
        analyst = _user(cur, "analyst", "ra_analyst")
        other = _user(cur, "analyst", "ra_other")
        admin = _user(cur, "admin", "ra_admin")
        viewer = _user(cur, "viewer", "ra_viewer")
    run = runs.submit_run(conn, user=analyst, report_id="search_results",
                          parameters={"criteria": {"text": word}})
    out = runs.execute_run(conn, run["id"])
    assert out["status"] == "completed", out
    yield {"conn": conn, "word": word, "run": runs.get_run(conn, run["id"], user=analyst),
           "analyst": analyst, "other": other, "admin": admin, "viewer": viewer,
           "pg_db": pg_db}
    conn.close()


def _create(world, fmt, dataset_key=None, user=None, run_id=None):
    user = user or world["analyst"]
    run_id = run_id or world["run"]["id"]
    plan = artifacts.request_artifact(world["conn"], run_id, user=user, fmt=fmt,
                                      dataset_key=dataset_key)
    if "existing" in plan:
        return {"status": "existing", "artifact": plan["existing"]}
    return artifacts.create_artifact(world["conn"], run_id=run_id, fmt=fmt,
                                     dataset_key=dataset_key, creator_id=user.id)


def _row(world, artifact_id):
    return artifacts.artifact_for_download(world["conn"], artifact_id, user=world["analyst"])


def test_every_format_is_stored_with_its_digest_and_manifest(world):
    run = world["run"]
    for fmt, key in (("json", None), ("html", None), ("xlsx", None),
                     ("csv", "search_results.matches@1")):
        out = _create(world, fmt, key)
        assert out["status"] in ("created", "existing"), out
        row = _row(world, out["artifact"]["id"])
        content = row["content"]
        assert row["sha256"] == hashlib.sha256(content).hexdigest()
        assert row["byte_size"] == len(content)
        m = row["manifest"]
        assert m["manifest_version"] == "report-manifest/2"
        assert m["artifact"] == dict(m["artifact"], format=fmt, sha256=row["sha256"],
                                     bytes=len(content),
                                     renderer_version=renderers.RENDERERS[fmt])
        assert m["report"]["key"] == "search_results@1"
        assert m["report"]["definition_fingerprint"] == run["definition_fingerprint"]
        assert m["run"]["id"] == run["id"]
        assert m["run"]["criteria"] == run["parameters"]["criteria"]
        assert m["run"]["criteria_fingerprint"] == run["criteria_fingerprint"]
        assert m["run"]["parameters_fingerprint"] == run["parameters_fingerprint"]
        assert m["snapshot"] == {"id": run["snapshot"], "taken_at": run["snapshot_at"],
                                 "isolation": runs.ISOLATION}
        assert [d["query_fingerprint"] for d in m["datasets"]] == \
            [d["query_fingerprint"] for d in run["datasets"]]
        assert m["generated"]["by"]["username"] == world["analyst"].username
        assert m["artifact"]["notes"]["null_representation"]
        expected_rows = 3 if key else 4          # 3 matches (+1 count row)
        assert m["row_count"] == expected_rows and m["truncated"] is False
        assert row["manifest_sha256"] == hashlib.sha256(
            renderers.canonical_json(m)).hexdigest()
        assert artifacts.verify(row)["ok"]


def test_rendering_is_deterministic_and_requests_are_idempotent(world):
    conn = world["conn"]
    with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT " + runs._RUN_COLUMNS + " FROM report_runs WHERE id = %s",
                    (world["run"]["id"],))
        run_row = cur.fetchone()
    doc = artifacts.run_document(conn, run_row)
    first = {f: renderers.render(doc, f, "search_results.matches@1" if f == "csv" else None)
             for f in renderers.RENDERERS}
    time.sleep(2.1)          # ZIP timestamps have 2 s resolution: prove none leak in
    doc_again = artifacts.run_document(conn, run_row)
    for fmt, rendering in first.items():
        again = renderers.render(doc_again, fmt,
                                 "search_results.matches@1" if fmt == "csv" else None)
        assert again.content == rendering.content, fmt
    a = _create(world, "json")
    b = _create(world, "json", user=world["admin"])
    assert a["artifact"]["id"] == b["artifact"]["id"] and b["status"] == "existing"
    stored = _row(world, a["artifact"]["id"])["content"]
    assert stored == first["json"].content


def test_the_database_refuses_changes_and_mismatched_digests(world):
    conn = world["conn"]
    aid = _create(world, "json")["artifact"]["id"]
    conn.rollback()
    for statement in ("UPDATE report_artifacts SET filename = 'x.json' WHERE id = %s",
                      "UPDATE report_artifacts SET content = 'x'::bytea,"
                      " sha256 = encode(sha256('x'::bytea), 'hex'), byte_size = 1 WHERE id = %s"):
        with pytest.raises(psycopg2.errors.IntegrityConstraintViolation, match="immutable"):
            with conn.cursor() as cur:
                cur.execute(statement, (aid,))
        conn.rollback()
    base = ("INSERT INTO report_artifacts (run_id, format, renderer_version, filename,"
            " media_type, byte_size, sha256, content, manifest, manifest_sha256,"
            " creator_username, creator_role) VALUES (%s, 'json', 'test/9', %s,"
            " 'application/json', %s, %s, %s, '{}'::jsonb, %s, 'x', 'analyst')")
    good = hashlib.sha256(b"abc").hexdigest()
    cases = {"ck_report_artifacts_digest": ("ok.json", 3, "0" * 64),
             "ck_report_artifacts_size": ("ok.json", 4, good),
             "ck_report_artifacts_filename": ("../etc/passwd", 3, good)}
    for constraint, (name, size, digest) in cases.items():
        with pytest.raises(psycopg2.errors.CheckViolation, match=constraint):
            with conn.cursor() as cur:
                cur.execute(base, (world["run"]["id"], name, size, digest,
                                   psycopg2.Binary(b"abc"), "0" * 64))
        conn.rollback()


def test_deleting_the_creator_keeps_the_artifact(world):
    conn = world["conn"]
    with conn, conn.cursor() as cur:
        temp = _user(cur, "admin", f"ra_temp{uuid.uuid4().hex[:4]}")
    out = _create(world, "csv", "search_results.count@1", user=temp)
    assert out["status"] == "created"
    with conn, conn.cursor() as cur:
        cur.execute("DELETE FROM users WHERE id = %s", (temp.id,))
        cur.execute("SELECT created_by, creator_username FROM report_artifacts WHERE id = %s",
                    (out["artifact"]["id"],))
        assert cur.fetchone() == (None, temp.username)


def test_a_tampered_manifest_is_detected_and_not_sent(world):
    conn = world["conn"]
    aid = _create(world, "html")["artifact"]["id"]
    conn.rollback()
    with conn, conn.cursor() as cur:       # a superuser bypassing the trigger
        cur.execute("ALTER TABLE report_artifacts DISABLE TRIGGER trg_report_artifacts_immutable")
        cur.execute("UPDATE report_artifacts SET manifest = jsonb_set(manifest,"
                    " '{row_count}', '999') WHERE id = %s", (aid,))
        cur.execute("ALTER TABLE report_artifacts ENABLE TRIGGER trg_report_artifacts_immutable")
    try:
        result = artifacts.verify_artifact(conn, aid, user=world["analyst"])
        assert result["ok"] is False and result["checks"]["manifest_sha256"] is False
        assert result["checks"]["content_sha256"] is True
        with pytest.raises(runs.ReportRunError) as exc:
            artifacts.artifact_for_download(conn, aid, user=world["analyst"])
        assert exc.value.code == "INTEGRITY_FAILED" and exc.value.status == 500
    finally:
        with conn, conn.cursor() as cur:
            cur.execute("DELETE FROM report_artifacts WHERE id = %s", (aid,))


def test_refusals(world):
    conn = world["conn"]
    run_id = world["run"]["id"]

    def refused(code, status, **kw):
        args = dict(user=world["analyst"], fmt="json", dataset_key=None)
        args.update(kw)
        with pytest.raises(runs.ReportRunError) as exc:
            artifacts.request_artifact(conn, kw.pop("run_id", run_id), **{
                k: v for k, v in args.items() if k != "run_id"})
        assert (exc.value.code, exc.value.status) == (code, status), exc.value.message
        return exc.value.message

    assert "step 18" in refused("VALIDATION_FAILED", 400, fmt="pdf")
    refused("VALIDATION_FAILED", 400, fmt="docx")
    refused("VALIDATION_FAILED", 400, fmt="csv")
    refused("VALIDATION_FAILED", 400, fmt="csv", dataset_key="nope@1")
    refused("VALIDATION_FAILED", 400, fmt="json", dataset_key="search_results.count@1")
    refused("NOT_FOUND", 404, user=world["other"])
    refused("FORBIDDEN", 403, user=world["viewer"])
    queued = runs.submit_run(conn, user=world["analyst"], report_id="search_results",
                             parameters={"criteria": {"text": world["word"]}})
    with pytest.raises(runs.ReportRunError) as exc:
        artifacts.request_artifact(conn, queued["id"], user=world["analyst"], fmt="json")
    assert (exc.value.code, exc.value.status) == ("RUN_NOT_COMPLETED", 409)
    # the administrator may render anyone's run; the creator is re-read at render time
    assert _create(world, "csv", "search_results.count@1", user=world["admin"])["status"] \
        in ("created", "existing")
    with conn, conn.cursor() as cur:
        cur.execute("UPDATE users SET is_active = FALSE WHERE id = %s", (world["other"].id,))
    out = artifacts.create_artifact(conn, run_id=run_id, fmt="html", dataset_key=None,
                                    creator_id=world["other"].id)
    assert out == {"status": "refused", "reason": "creator_not_permitted"}
    with conn, conn.cursor() as cur:
        cur.execute("UPDATE users SET is_active = TRUE WHERE id = %s", (world["other"].id,))
    out = artifacts.create_artifact(conn, run_id=run_id, fmt="xlsx", dataset_key=None,
                                    creator_id=world["other"].id)   # active, not the owner
    assert out["status"] == "refused" and out["reason"] == "not_found"


def test_hostile_and_non_latin_text_survives_every_format(world):
    json_body = json.loads(_row(world, _create(world, "json")["artifact"]["id"])["content"])
    names = {r["file_name"] for r in json_body["datasets"][0]["rows"]}
    assert f"{HOSTILE}.pdf" in names and f"{ARABIC}.pdf" in names

    html_body = _row(world, _create(world, "html")["artifact"]["id"])["content"].decode()
    assert "<script>alert(1)</script>" not in html_body
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html_body
    assert f"<bdi>{ARABIC}.pdf</bdi>" in html_body
    assert "default-src 'none'" in html_body

    csv_bytes = _row(world, _create(world, "csv", "search_results.matches@1")
                     ["artifact"]["id"])["content"]
    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
    by_name = {r["file_name"] for r in rows}
    assert "'" + f"{HOSTILE}.pdf" in by_name and f"{ARABIC}.pdf" in by_name

    from openpyxl import load_workbook

    xlsx = _row(world, _create(world, "xlsx")["artifact"]["id"])
    book = load_workbook(io.BytesIO(xlsx["content"]))
    assert book.sheetnames[0] == "Provenance"
    listing = book[book.sheetnames[1]]
    values = {c.value for c in listing["B"]}
    assert "'" + f"{HOSTILE}.pdf" in values and f"{ARABIC}.pdf" in values
    assert xlsx["manifest"]["artifact"]["notes"]["formula_guarded_cells"] >= 1
    with zipfile.ZipFile(io.BytesIO(xlsx["content"])) as z:
        pinned = {i.date_time for i in z.infolist()}
    assert len(pinned) == 1


def test_truncation_is_stated_in_the_content_and_the_manifest(world):
    import dataclasses

    from core.reporting.registry import ReportRegistry

    matches = dataclasses.replace(REGISTRY.dataset("search_results.matches@1"), row_limit=2)
    registry = ReportRegistry(reports=REGISTRY.reports,
                              datasets=tuple(matches if d.key == matches.key else d
                                             for d in REGISTRY.datasets),
                              help_topics=REGISTRY.help_topics)
    conn = world["conn"]
    run = runs.submit_run(conn, user=world["analyst"], report_id="search_results",
                          parameters={"criteria": {"text": world["word"]}}, registry=registry)
    assert runs.execute_run(conn, run["id"], registry=registry)["status"] == "completed"
    with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT " + runs._RUN_COLUMNS + " FROM report_runs WHERE id = %s",
                    (run["id"],))
        run_row = cur.fetchone()
    doc = artifacts.run_document(conn, run_row, registry)
    conn.rollback()
    html_body = renderers.render(doc, "html").content.decode()
    assert "shortened: the first 2 rows are included; more rows existed (capped at 2)" in html_body
    assert "complete: 1 rows (exact)" in html_body
    rendering = renderers.render(doc, "csv", "search_results.matches@1")
    assert len(list(csv.reader(io.StringIO(rendering.content.decode("utf-8-sig"))))) == 3
    manifest = artifacts.build_manifest(doc, rendering, content_sha256="0" * 64, byte_size=1,
                                        creator={}, generated_at="t", job_id=None)
    assert manifest["truncated"] is True and manifest["row_count"] == 2
