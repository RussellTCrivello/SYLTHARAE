"""The DOCX reader must return every paragraph and table, once, in order.

Regression test for DOCX-DROP-01: the reader matched body elements to
python-docx proxies and remembered the ones it had used by ``id()``. Those
proxies are rebuilt on every ``doc.paragraphs`` access and freed straight
away, so CPython reused their addresses and a *different* paragraph could
look "already processed" and be skipped. Nothing reported it; the text was
simply absent from storage, search and display. On a 200-section document
the old reader lost 54-73 of 400 paragraphs per read, a different set each
time, which is why a four-paragraph integration fixture failed only now and
then.
"""

import logging

import pytest

from reader_file.readers.read_office import OfficeFileReader

SECTIONS = 200


@pytest.fixture(scope="module")
def long_docx(tmp_path_factory):
    from docx import Document

    doc = Document()
    expected = []  # ("paragraph", text) and ("table", first cell), body order
    for i in range(SECTIONS):
        doc.add_heading(f"Section {i}", level=1 + i % 3)
        expected.append(("paragraph", f"Section {i}"))
        doc.add_paragraph(f"Body text {i}.")
        expected.append(("paragraph", f"Body text {i}."))
        if i % 5 == 0:
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = f"key {i}"
            table.cell(0, 1).text = "value"
            expected.append(("table", f"key {i}"))
    path = tmp_path_factory.mktemp("docx_completeness") / "long.docx"
    doc.save(path)
    return str(path), expected


@pytest.fixture(scope="module")
def reader():
    return OfficeFileReader()


def _body(result):
    """The reader's elements as (type, identifying text), dropping blank paragraphs."""
    out = []
    for element in result["elements"]:
        if element["type"] == "paragraph" and element["text"]:
            out.append(("paragraph", element["text"]))
        elif element["type"] == "table":
            out.append(("table", element["rows"][0][0]))
    return out


@pytest.mark.parametrize("attempt", range(3))
def test_every_paragraph_and_table_is_returned_once_in_order(reader, long_docx, attempt):
    # Repeated reads: the old defect lost a different set of paragraphs each time.
    path, expected = long_docx
    got = _body(reader.read_docx_file(path))
    missing = [item for item in expected if item not in got]
    assert not missing, f"{len(missing)} of {len(expected)} body elements dropped, e.g. {missing[:5]}"
    assert got == expected


def test_positions_follow_document_order_without_gaps(reader, long_docx):
    path, _ = long_docx
    positions = [e["position"] for e in reader.read_docx_file(path)["elements"]]
    assert positions == list(range(len(positions)))


def test_backward_compatible_paragraph_list_is_complete(reader, long_docx):
    path, expected = long_docx
    result = reader.read_docx_file(path)
    texts = [p["text"] for p in result["paragraphs"]]
    assert texts == [text for kind, text in expected if kind == "paragraph"]
    styles = {p["text"]: p["style"] for p in result["paragraphs"]}
    assert styles["Section 0"] == "Heading 1" and styles["Section 1"] == "Heading 2"
    assert styles["Body text 0."] == "Normal"


def test_reading_does_not_log_errors(reader, long_docx, caplog):
    path, _ = long_docx
    with caplog.at_level(logging.WARNING):
        reader.read_docx_file(path)
    assert not [r for r in caplog.records if r.levelno >= logging.ERROR]
